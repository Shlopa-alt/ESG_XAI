"""Срезы Excel для общих методов: листы-таблицы, строка заголовка, буквенные рейтинги ESG-…

Строки — из DS-05 (ESG.xlsx: рэнкинг RAEX за 01.2026 и рейтинги «Эксперт РА»).
"""

from pathlib import Path

import pytest
from openpyxl import Workbook

from app.services.data_files import is_table, table_columns
from app.services.run_templates import TEMPLATES, ParamError, check_params


@pytest.fixture
def book(tmp_path: Path) -> Path:
    wb = Workbook()
    notes = wb.active
    notes.title = "Обозначения"
    notes["B3"] = "E-рэнк (E-рейтинг) — это показатель в системе ESG-оценки бизнеса, который отражает уровень "
    rank = wb.create_sheet("01.2026")
    rank.append(["№", "Название", "Страна", "MOEX", "Подотрасль", "Е - рэнк", "S - рэнк", "G - рэнк"])
    rank.append([1, "Норильский никель", "Россия", "GMKN", "Горнодобывающая промышленность и металлургия", 2, 6, 6])
    rank.append([2, "ФосАгро", "Россия", "PHOR", "Химическая промышленность", 6, 1, 16, None,
                 "https://raex-rr.com/ESG/ESG_companies/ESG_rating_companies/2024.12/"])  # заметка без заголовка
    rank.append([3, "Сбербанк", "Россия", "SBER", "Банки", 1, 8, 9])
    expert = wb.create_sheet("эксперт ра")
    expert.append(["Объект", "Рейтинг", "Прогноз", "Обновлен"])
    for row in (['ПАО "РУСГИДРО"', "ESG-AA", "Стабильный"], ['ПАО "СПБ БИРЖА"', "ESG-BBB", "Стабильный"],
                ['ООО "ЛЕГЕНДА"', "ESG-A–", "Стабильный"], ['ПАО "АК БАРС" БАНК', "ESG-BBB+", "Стабильный"],
                ['ПАО "МТС"', "ESG-AA", "Стабильный"], ['ПАО "ГК "САМОЛЕТ"', "отозван", "—"]):
        expert.append(row)
    path = tmp_path / "ESG.xlsx"
    wb.save(path)
    return path


def test_sheets_and_header(book: Path) -> None:
    t = table_columns(book)
    assert t.sheets == ["01.2026", "эксперт ра"]  # «Обозначения» — текст, не таблица
    assert t.sheet == "01.2026" and t.columns[-3:] == ["Е - рэнк", "S - рэнк", "G - рэнк"]
    assert len(t.columns) == 8  # заметка справа без заголовка отброшена
    assert {"Е - рэнк", "S - рэнк", "G - рэнк"} <= t.numeric and "MOEX" not in t.numeric


def test_esg_letter_ratings_are_numeric(book: Path) -> None:
    t = table_columns(book, "эксперт ра")
    assert t.numeric == {"Рейтинг"}
    with pytest.raises(ValueError):
        table_columns(book, "Обозначения")


def test_sheet_param_checked(book: Path) -> None:
    rank = TEMPLATES["rank-agreement"]
    t = table_columns(book, "01.2026")
    ok = check_params(rank, {"columns": ["Е - рэнк", "G - рэнк"]}, t.columns, t.numeric, t.sheets, t.sheet)
    assert ok["sheet"] == "01.2026"  # пусто — первый лист-таблица, записан явно
    with pytest.raises(ParamError):
        check_params(rank, {"sheet": "Обозначения", "columns": ["Е - рэнк", "G - рэнк"]}, t.columns, t.numeric,
                     t.sheets, t.sheet)
    with pytest.raises(ParamError):  # у CSV листов нет
        check_params(rank, {"sheet": "01.2026", "columns": ["a", "b"]}, ["a", "b"], {"a", "b"}, [], "")
    assert is_table("C:/x/ESG.xlsx") and is_table("a.csv") and not is_table("a.pdf")
