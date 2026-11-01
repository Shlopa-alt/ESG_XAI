"""Файлы срезов на сервере: путь хоста ↔ путь в контейнере, листы и столбцы таблицы.

Срезы записаны в платформе путями хоста (ESG_DATA_HOST), в контейнере та же папка — ESG_DATA_DIR.
Таблица — CSV/TSV или лист книги Excel (.xlsx). В книге лист-таблица — тот, где в первых пяти
строках есть строка заголовка: не меньше трёх непустых ячеек, все — короткий текст; столбцы без
имени (заметки справа) отбрасываются. Так же читают методы: runner/methods/common.py.
"""

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from app.config import Settings

# буквы рейтингов — порядковая шкала, методы переводят их в числа (ESG-BBB+, ruAA- — тоже)
RATING_SCALE = {"AAA", "AA", "A", "BBB", "BB", "B", "CCC", "CC", "C", "D"}
RATING = re.compile(r"(?:ESG-|RU)?(AAA|AA|A|BBB|BB|B|CCC|CC|C|D)[+\-–−]?")
EXCEL = (".xlsx", ".xlsm")
TEXT = (".csv", ".tsv", ".txt")
HEADER_MAX = 60  # длиннее — это заголовок таблицы или пояснение, а не имя столбца
SAMPLE = 200
# пометки «значения нет» — как пустая ячейка (у «Эксперт РА» четверть рейтингов — «отозван»);
# методы их тоже читают как пропуск
MISSING = {"", "-", "–", "—", "−", "н/д", "нд", "нет", "n/a", "na", "nan", "отозван", "отозвано", "отозвана"}
NUMBER = re.compile(r"[-+]?(\d+\.?\d*|\.\d+)([eE][-+]?\d+)?")


def to_container(settings: Settings, host_path: str) -> Path | None:
    p = host_path.replace("\\", "/")
    host = settings.esg_data_host.replace("\\", "/").rstrip("/")
    if host and p.lower().startswith(host.lower() + "/"):
        return Path(settings.esg_data_dir) / p[len(host) + 1:]
    if p.startswith(settings.esg_data_dir.rstrip("/") + "/"):
        return Path(p)
    return None




@dataclass
class Table:
    columns: list[str]
    numeric: set[str]
    sheets: list[str] = field(default_factory=list)  # листы-таблицы книги; у CSV — пусто
    sheet: str = ""  # какой лист прочитан


def is_table(path: str | Path) -> bool:
    return str(path).lower().endswith(TEXT + EXCEL)


def _is_number(v: str) -> bool:
    v = v.strip()
    return bool(NUMBER.fullmatch(v.replace(",", ".")) or RATING.fullmatch(v.upper()))


def _numeric(columns: list[str], rows: list[list[Any]]) -> set[str]:
    """Числовой столбец — не меньше 80 % непустых значений числа или буквы рейтинга."""
    out = set()
    for j, c in enumerate(columns):
        vals = [r[j] for r in rows if j < len(r) and r[j] not in (None, "")
                and str(r[j]).strip().lower() not in MISSING]
        good = sum(1 for v in vals if (isinstance(v, int | float) and not isinstance(v, bool))
                   or (isinstance(v, str) and _is_number(v)))
        if vals and good >= 0.8 * len(vals):
            out.add(c)
    return out


def header_row(rows: list[tuple[Any, ...]]) -> int | None:
    for i, row in enumerate(rows[:5]):
        cells = [c for c in row if c not in (None, "") and str(c).strip()]
        if len(cells) >= 3 and all(isinstance(c, str) and len(c.strip()) <= HEADER_MAX for c in cells):
            return i
    return None


def sheet_table(rows: list[tuple[Any, ...]]) -> tuple[list[int], list[str], list[list[Any]]] | None:
    """Лист как таблица: номера столбцов, имена (повтор — «имя (2)»), строки данных; не таблица — None."""
    h = header_row(rows)
    if h is None:
        return None
    idx, names = [], []
    for j, c in enumerate(rows[h]):
        name = " ".join(str(c).split()) if c not in (None, "") else ""
        if not name:
            continue
        n, k = name, 2
        while n in names:
            n, k = f"{name} ({k})", k + 1
        idx.append(j)
        names.append(n)
    data = [[r[j] if j < len(r) else None for j in idx] for r in rows[h + 1:]]
    data = [r for r in data if any(v not in (None, "") and str(v).strip() for v in r)]
    return (idx, names, data) if data else None


def excel_sheets(path: Path, limit: int | None = SAMPLE) -> dict[str, tuple[list[str], list[list[Any]]]]:
    """Листы-таблицы книги: {лист: (столбцы, строки)} — строк не больше limit (None — все)."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    out = {}
    try:
        for ws in wb.worksheets:
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if limit is not None and i > limit + 5:
                    break
                rows.append(tuple(row))
            t = sheet_table(rows)
            if t is not None:
                out[ws.title] = (t[1], t[2][:limit] if limit is not None else t[2])
    finally:
        wb.close()
    return out


def table_columns(path: Path, sheet: str = "") -> Table:
    """Столбцы таблицы и какие из них числовые (по первым строкам; пустые значения не мешают,
    буквенная шкала рейтингов AAA…D, ESG-BBB+ считается числовой). У книги — лист sheet (пусто —
    первый лист-таблица). Не таблица или нет такого листа — ValueError."""
    suffix = path.suffix.lower()
    if suffix in EXCEL:
        sheets = excel_sheets(path)
        if not sheets:
            raise ValueError("В книге нет листа-таблицы со строкой заголовка")
        name = sheet or next(iter(sheets))
        if name not in sheets:
            raise ValueError(f"В книге нет листа-таблицы «{name}»; есть: {', '.join(sheets)}")
        columns, rows = sheets[name]
        return Table(columns, _numeric(columns, rows), list(sheets), name)
    if suffix not in TEXT:
        raise ValueError(f"Общие методы читают CSV и Excel (.xlsx), а срез — {path.suffix or 'без расширения'}")
    if sheet:
        raise ValueError("Лист выбирают только у книги Excel")
    with path.open(encoding="utf-8-sig", errors="replace", newline="") as f:
        first = f.readline()
        f.seek(0)
        try:
            dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(first, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(f, dialect)
        columns = [c.strip() for c in next(reader, [])]
        rows = [row for _, row in zip(range(SAMPLE), reader, strict=False)]
    return Table(columns, _numeric(columns, [r[: len(columns)] for r in rows]))
