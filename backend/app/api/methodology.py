import io
from datetime import datetime
from typing import Annotated, Literal

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.deps import current_user, require_role
from app.db import get_db
from app.models import (
    ChangeRequest,
    CrBasis,
    Experiment,
    ExperimentReport,
    KbRef,
    MethodologySection,
    MethodologyVersion,
    Review,
    User,
)
from app.services import methodology as svc
from app.services import qwen_cr
from app.services.codes import next_code
from app.services.experiments import DomainError

router = APIRouter(tags=["methodology"])

Researcher = Depends(require_role("researcher"))
ReviewerRole = Depends(require_role("reviewer"))


class SectionOut(BaseModel):
    number: str
    title: str
    content: str
    stability: str
    kind: str  # direction | variant | task | section
    parent: str  # номер родителя; пусто — корень методики
    position: int


class VersionOut(BaseModel):
    version: str
    notes: str
    created_at: datetime


class MethodologyOut(BaseModel):
    version: str | None
    sections: list[SectionOut]


class BasisOut(BaseModel):
    basis_type: str
    ref: str
    valid: bool


class CrReviewOut(BaseModel):
    reviewer: str
    verdict: str
    text: str
    created_at: datetime


class CrOut(BaseModel):
    code: str
    section: str  # edit — изменяемый узел, add — родитель нового (пусто — корень)
    section_title: str
    action: str
    new_kind: str
    new_title: str
    title: str
    description: str
    proposed_content: str
    status: str
    target_version: str
    bases: list[BasisOut]
    reviews: list[CrReviewOut]
    blockers: list[str]
    source: str  # "" — люди, "qwen" — черновик Qwen по отчёту
    report: str  # отчёт-источник: «EXP-001 v1»
    report_experiment: str
    problems: str  # что проверить в черновике Qwen
    author: str
    editable: bool  # текст ещё можно править: CR открыт и одобряющей рецензии нет


class CrIn(BaseModel):
    code: str | None = Field(default=None, pattern=r"^CR-\d{2,3}$")  # нет — следующий свободный
    section: str = ""  # edit — номер узла; add — номер родителя, пусто — корень методики
    action: Literal["edit", "add"] = "edit"
    new_kind: Literal["direction", "variant", "task", "section", ""] = ""
    new_title: str = Field(default="", max_length=512)
    title: str = Field(min_length=3)
    description: str = ""
    proposed_content: str = ""


class CrPatch(BaseModel):
    status: Literal["review", "accepted", "rejected"] | None = None
    # правка текста (researcher+), пока CR открыт и одобряющей рецензии нет — иначе одобряли бы не то
    title: str | None = Field(default=None, min_length=3, max_length=512)
    description: str | None = None
    proposed_content: str | None = None
    new_title: str | None = Field(default=None, min_length=3, max_length=512)


class BasisIn(BaseModel):
    basis_type: Literal["experiment", "kb_fact", "kb_article"]
    experiment_code: str | None = None
    kbrag_id: str = ""
    kb_path: str = ""
    title: str = ""
    doi: str = ""
    quote: str = ""
    page: int | None = None


class CrReviewIn(BaseModel):
    verdict: Literal["approve", "changes", "reject"]
    text: str = ""


def _section(db: Session, number: str) -> MethodologySection:
    s = db.scalar(select(MethodologySection).where(MethodologySection.number == number))
    if s is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, f"Нет раздела {number}")
    return s


def _cr(db: Session, code: str) -> ChangeRequest:
    cr = db.scalar(select(ChangeRequest).where(ChangeRequest.code == code))
    if cr is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, f"Нет предложения {code}")
    return cr


def _cr_out(db: Session, cr: ChangeRequest) -> CrOut:
    section = db.get(MethodologySection, cr.section_id) if cr.section_id else None
    valid = set(svc.valid_bases(db, cr))
    bases = []
    for b in db.scalars(select(CrBasis).where(CrBasis.cr_id == cr.id)):
        if b.experiment_id is not None:
            exp = db.get(Experiment, b.experiment_id)
            label = f"эксперимент {exp.code}" if exp else "эксперимент ?"
            ref = exp.code if exp else "?"
        else:
            kb = db.get(KbRef, b.kb_ref_id) if b.kb_ref_id else None
            ref = (kb.doi or kb.kbrag_id) if kb else "?"
            label = (f"факт KBRAG {kb.kbrag_id}" if b.basis_type == "kb_fact"
                     else f"статья KBRAG {kb.doi}") if kb else "?"
        bases.append(BasisOut(basis_type=b.basis_type, ref=ref, valid=label in valid))
    reviews = db.execute(
        select(Review, User.login).join(User, User.id == Review.reviewer_id)
        .where(Review.target_type == "cr", Review.target_id == cr.id).order_by(Review.created_at)
    ).all()
    report = db.get(ExperimentReport, cr.report_id) if cr.report_id else None
    rep_exp = db.get(Experiment, report.experiment_id) if report else None
    author = db.get(User, cr.author_id) if cr.author_id else None
    return CrOut(
        code=cr.code,
        section=section.number if section else "",
        section_title=section.title if section else "корень методики",
        action=cr.action,
        new_kind=cr.new_kind,
        new_title=cr.new_title,
        title=cr.title,
        description=cr.description,
        proposed_content=cr.proposed_content,
        status=cr.status,
        target_version=cr.target_version,
        bases=bases,
        reviews=[CrReviewOut(reviewer=login, verdict=r.verdict, text=r.text,
                             created_at=r.created_at) for r, login in reviews],
        blockers=[] if cr.status in ("accepted", "rejected") else svc.why_not_acceptable(db, cr),
        source=cr.source,
        report=f"{rep_exp.code} v{report.version}" if report and rep_exp else "",
        report_experiment=rep_exp.code if rep_exp else "",
        problems=cr.problems if cr.status in ("proposed", "review") else "",
        author=author.login if author else ("Qwen" if cr.source == "qwen" else ""),
        editable=cr.status in ("proposed", "review") and not svc.has_approval(db, cr),
    )


@router.get("/methodology", response_model=MethodologyOut, dependencies=[Depends(current_user)])
def methodology(db: Annotated[Session, Depends(get_db)]) -> MethodologyOut:
    cur = svc.current_version(db)
    return MethodologyOut(version=cur.version if cur else None, sections=_tree(db))


def _tree(db: Session) -> list[SectionOut]:
    """Узлы в порядке обхода дерева: родитель, затем его разделы, затем дочерние узлы."""
    nodes = list(db.scalars(select(MethodologySection)))
    number = {n.id: n.number for n in nodes}
    children: dict[int | None, list[MethodologySection]] = {}
    for n in nodes:
        children.setdefault(n.parent_id, []).append(n)
    order = {"section": 0, "direction": 1, "variant": 1, "task": 1}
    out: list[SectionOut] = []

    def walk(pid: int | None) -> None:
        for n in sorted(children.get(pid, []), key=lambda x: (order.get(x.kind, 1), x.position, x.number)):
            out.append(SectionOut(number=n.number, title=n.title, content=n.content, stability=n.stability,
                                  kind=n.kind, parent=number.get(n.parent_id, "") if n.parent_id else "",
                                  position=n.position))
            walk(n.id)

    walk(None)
    return out


@router.get("/methodology/versions", response_model=list[VersionOut],
            dependencies=[Depends(current_user)])
def versions(db: Annotated[Session, Depends(get_db)]) -> list[VersionOut]:
    rows = db.scalars(select(MethodologyVersion).order_by(MethodologyVersion.id.desc()))
    return [VersionOut(version=v.version, notes=v.notes, created_at=v.created_at) for v in rows]


def _version(db: Session, v: str) -> MethodologyVersion:
    found = db.scalar(select(MethodologyVersion).where(MethodologyVersion.version == v))
    if found is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, f"Нет версии {v}")
    return found


@router.get("/methodology/diff", dependencies=[Depends(current_user)])
def methodology_diff(
    db: Annotated[Session, Depends(get_db)],
    from_: Annotated[str, Query(alias="from")],
    to: str,
) -> list[dict[str, object]]:
    return svc.diff(_version(db, from_).sections_content, _version(db, to).sections_content)


@router.get("/methodology/versions/{v}", dependencies=[Depends(current_user)])
def version(v: str, db: Annotated[Session, Depends(get_db)]) -> dict[str, object]:
    found = _version(db, v)
    return {"version": found.version, "notes": found.notes, "sections": found.sections_content}


@router.get("/methodology/export.docx", dependencies=[Depends(current_user)])
def export_docx(db: Annotated[Session, Depends(get_db)]) -> StreamingResponse:
    from docx import Document  # type: ignore[import-untyped]  # python-docx

    cur = svc.current_version(db)
    doc = Document()
    doc.add_heading(f"Методология ESG Open Research, версия {cur.version if cur else '—'}", 0)
    labels = {"stable": "устойчив", "draft": "черновик", "contested": "оспаривается"}
    depth = {"": 0}
    for s in _tree(db):
        depth[s.number] = depth.get(s.parent, 0) + 1
        doc.add_heading(f"{s.number}. {s.title}", level=min(depth[s.number], 4))
        doc.add_paragraph(f"Устойчивость: {labels.get(s.stability, s.stability)}")
        doc.add_paragraph(s.content)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    name = f"methodology-{cur.version if cur else 'draft'}.docx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )


@router.get("/change-requests", response_model=list[CrOut], dependencies=[Depends(current_user)])
def list_crs(db: Annotated[Session, Depends(get_db)]) -> list[CrOut]:
    return [_cr_out(db, cr) for cr in db.scalars(select(ChangeRequest).order_by(ChangeRequest.code))]


@router.post("/change-requests", response_model=CrOut, status_code=status.HTTP_201_CREATED)
def create_cr(
    body: CrIn, db: Annotated[Session, Depends(get_db)], user: Annotated[User, Researcher]
) -> CrOut:
    code = body.code or next_code(db, ChangeRequest.code, "CR", 2)
    if db.scalar(select(ChangeRequest).where(ChangeRequest.code == code)):
        raise HTTPException(status.HTTP_409_CONFLICT, f"Код {code} уже занят")
    section = _section(db, body.section) if body.section else None
    if body.action == "edit":
        if section is None:
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Какой узел меняется — section")
    else:
        allowed = svc.CHILD_KINDS[section.kind if section else None]
        if body.new_kind not in allowed:
            where = f"под «{section.number}»" if section else "в корень методики"
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY,
                                f"Добавить {where} можно: {', '.join(allowed) or 'ничего'}")
        if len(body.new_title.strip()) < 3:
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Заголовок нового узла — new_title")
        if not body.proposed_content.strip():
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Текст нового узла — proposed_content")
    cr = ChangeRequest(code=code, section_id=section.id if section else None, action=body.action,
                       new_kind=body.new_kind if body.action == "add" else "",
                       new_title=body.new_title.strip() if body.action == "add" else "",
                       title=body.title, description=body.description,
                       proposed_content=body.proposed_content, status="proposed", author_id=user.id)
    db.add(cr)
    db.commit()
    return _cr_out(db, cr)


@router.post("/change-requests/{code}/bases", response_model=CrOut, dependencies=[Researcher])
def add_basis(code: str, body: BasisIn, db: Annotated[Session, Depends(get_db)]) -> CrOut:
    cr = _cr(db, code)
    if body.basis_type == "experiment":
        exp = db.scalar(select(Experiment).where(Experiment.code == (body.experiment_code or "")))
        if exp is None:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "Нет такого эксперимента")
        db.add(CrBasis(cr_id=cr.id, basis_type="experiment", experiment_id=exp.id))
    else:
        if not body.kbrag_id:
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Нужен kbrag_id")
        ref = KbRef(kind="fact" if body.basis_type == "kb_fact" else "article",
                    kbrag_id=body.kbrag_id, kb_path=body.kb_path, title=body.title,
                    doi=body.doi, quote=body.quote, page=body.page)
        db.add(ref)
        db.flush()
        db.add(CrBasis(cr_id=cr.id, basis_type=body.basis_type, kb_ref_id=ref.id))
    db.commit()
    return _cr_out(db, cr)


@router.post("/change-requests/{code}/reviews", response_model=CrOut,
             status_code=status.HTTP_201_CREATED)
def review_cr(
    code: str, body: CrReviewIn, db: Annotated[Session, Depends(get_db)],
    user: Annotated[User, ReviewerRole],
) -> CrOut:
    cr = _cr(db, code)
    db.add(Review(target_type="cr", target_id=cr.id, reviewer_id=user.id,
                  verdict=body.verdict, text=body.text))
    if cr.status == "proposed":
        cr.status = "review"
    db.commit()
    return _cr_out(db, cr)


@router.patch("/change-requests/{code}", response_model=CrOut)
def patch_cr(
    code: str, body: CrPatch, db: Annotated[Session, Depends(get_db)],
    user: Annotated[User, Depends(current_user)],
) -> CrOut:
    cr = _cr(db, code)
    if cr.status in ("accepted", "rejected"):
        raise HTTPException(status.HTTP_409_CONFLICT, f"{cr.code} уже закрыт")
    text = body.model_dump(exclude_unset=True, exclude={"status"})
    if text:
        if user.role not in ("researcher", "owner"):
            raise HTTPException(status.HTTP_403_FORBIDDEN, "Править текст CR может researcher")
        if svc.has_approval(db, cr):
            raise HTTPException(status.HTTP_409_CONFLICT, "CR уже одобрен рецензентом — текст не меняется")
        if "new_title" in text and cr.action != "add":
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Заголовок нового узла — только у CR «добавить»")
        if cr.action == "add" and not (text.get("proposed_content", cr.proposed_content) or "").strip():
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Текст нового узла — proposed_content")
        for k, v in text.items():
            setattr(cr, k, (v or "").strip())
        cr.problems = qwen_cr.problems(db, cr)
    if body.status is None:
        db.commit()
        return _cr_out(db, cr)
    if body.status in ("accepted", "rejected") and user.role != "owner":
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Принимает и отклоняет CR только owner")
    if body.status == "review" and user.role not in ("researcher", "reviewer", "owner"):
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Недостаточно прав")
    if body.status == "accepted":
        try:
            svc.accept(db, cr, user.id)
        except DomainError as exc:
            db.rollback()
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, str(exc)) from exc
    else:
        cr.status = body.status
    db.commit()
    return _cr_out(db, cr)
