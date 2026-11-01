"""Общее для методов библиотеки (runner/methods): аргументы, чтение среза, run.json.

Каждый метод запускается одинаково (так его запускает исполнитель runner/worker):

    python <метод>.py <каталог датасета> <каталог результатов> --data <файл среза>
        --params '<JSON параметров>' --experiment EXP-0NN --dataset DS-0N

и пишет в <каталог результатов>/<время>/ результаты и run.json: коммит runner, sha256 среза,
sha256 окружения, seed, части, отчёт (ключевые цифры и графики для страницы эксперимента) —
правило 3. Ход и метрики печатаются строками `@event {...}`.
"""

import argparse
import hashlib
import json
import re
import subprocess
import sys
import time
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd

HERE = Path(__file__).resolve().parent
LOCK = HERE.parent / "hsu2025" / "requirements.lock"  # то же окружение, что у всех расчётов runner
RATING_SCALE = {"AAA": 10, "AA": 9, "A": 8, "BBB": 7, "BB": 6, "B": 5, "CCC": 4, "CC": 3, "C": 2, "D": 1}
# ESG-BBB+, ruAA-: приставка агентства отбрасывается, «+» и «−» — треть ступени шкалы
RATING = re.compile(r"(?:ESG-|RU)?(AAA|AA|A|BBB|BB|B|CCC|CC|C|D)([+\-–−]?)")
EXCEL = (".xlsx", ".xlsm")
HEADER_MAX = 60  # как app/services/data_files.py: длиннее — заголовок таблицы или пояснение


def event(kind: str, **data: object) -> None:
    """Событие для живого потока платформы; в консоли — просто строка."""
    print("@event " + json.dumps({"type": kind, **data}, ensure_ascii=False, default=str), flush=True)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(2**20), b""):
            h.update(block)
    return h.hexdigest()


def git_commit() -> str:
    try:
        commit = subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=HERE, text=True).strip()
        dirty = subprocess.check_output(["git", "status", "--porcelain", "--", "."], cwd=HERE, text=True).strip()
        return commit + ("-dirty" if dirty else "")
    except (OSError, subprocess.CalledProcessError):
        return ""


def parse_args(description: str) -> argparse.Namespace:
    ap = argparse.ArgumentParser(description=description)
    ap.add_argument("ds_dir", type=Path, help="каталог датасета")
    ap.add_argument("results_root", type=Path, help="каталог результатов")
    ap.add_argument("--data", type=Path, required=True, help="файл среза (CSV или книга Excel)")
    ap.add_argument("--params", type=json.loads, required=True, help="параметры метода, JSON")
    ap.add_argument("--experiment", required=True)
    ap.add_argument("--dataset", required=True)
    return ap.parse_args()


def _sheet_table(rows: list[tuple[Any, ...]]) -> pd.DataFrame | None:
    """Лист как таблица — как платформа (app/services/data_files.py): строка заголовка в первых пяти —
    не меньше трёх непустых ячеек, все — короткий текст; столбцы без имени отбрасываются, повтор
    имени — «имя (2)»; пустые строки — тоже."""
    head = None
    for i, row in enumerate(rows[:5]):
        cells = [c for c in row if c not in (None, "") and str(c).strip()]
        if len(cells) >= 3 and all(isinstance(c, str) and len(c.strip()) <= HEADER_MAX for c in cells):
            head = i
            break
    if head is None:
        return None
    idx, names = [], []
    for j, c in enumerate(rows[head]):
        name = " ".join(str(c).split()) if c not in (None, "") else ""
        if not name:
            continue
        n, k = name, 2
        while n in names:
            n, k = f"{name} ({k})", k + 1
        idx.append(j)
        names.append(n)
    data = [[r[j] if j < len(r) else None for j in idx] for r in rows[head + 1:]]
    data = [r for r in data if any(v not in (None, "") and str(v).strip() for v in r)]
    return pd.DataFrame(data, columns=names) if data else None


def load_table(path: Path, sheet: str = "") -> pd.DataFrame:
    """CSV/TSV или лист книги Excel (sheet; пусто — первый лист-таблица). У CSV разделитель — по
    строке заголовка (в данных бывают адреса с запятыми в кавычках, и угадывание по всему файлу
    ошибается). Как платформа: app/services/data_files.py."""
    if path.suffix.lower() in EXCEL:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            names = [ws.title for ws in wb.worksheets]
            for ws in wb.worksheets if not sheet else [wb[sheet]] if sheet in names else []:
                df = _sheet_table([tuple(r) for r in ws.iter_rows(values_only=True)])
                if df is not None:
                    return df
        finally:
            wb.close()
        raise SystemExit(f"В книге {path.name} нет листа-таблицы" + (f" «{sheet}»" if sheet else ""))
    with path.open(encoding="utf-8-sig", errors="replace") as f:
        header = f.readline()
    sep = max(("\t", ";", ","), key=header.count) if path.suffix.lower() != ".tsv" else "\t"
    return pd.read_csv(path, sep=sep, encoding="utf-8-sig")


def rating(value: object) -> float | None:
    """Буква рейтинга как число: AAA = 10 … D = 1; ESG-BBB+ = 7⅓, ruAA- = 8⅔; не рейтинг — None."""
    m = RATING.fullmatch(str(value).strip().upper())
    if m is None:
        return None
    return RATING_SCALE[m.group(1)] + {"+": 1 / 3, "": 0.0}.get(m.group(2), -1 / 3)


def numeric(df: pd.DataFrame, column: str) -> pd.Series:
    """Столбец как числа: буквенная шкала рейтингов (AAA…D, ESG-BBB+) — по порядку, остальное — to_numeric."""
    s = df[column]
    if not pd.api.types.is_numeric_dtype(s):  # в pandas 3 текст — тип str, а не object
        mapped = s.map(lambda v: None if pd.isna(v) else rating(v)).astype(float)
        if mapped.notna().sum() >= max(1, s.notna().sum() // 2):
            return mapped.astype(float)
        return pd.to_numeric(s.astype(str).str.replace(",", ".", regex=False), errors="coerce")
    return pd.to_numeric(s, errors="coerce")


def fmt(v: float, signed: bool = False) -> str:
    """Число по-русски: запятая, минус «−»."""
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return "—"
    text = f"{v:+.2f}" if signed else f"{v:.2f}"
    return text.replace(".", ",").replace("-", "−")


def write_run(args: argparse.Namespace, out: Path, part: str, started: str, t0: float, df: pd.DataFrame,
              seeds: dict[str, Any], summary: dict[str, Any], report: dict[str, Any], exit_code: int = 0) -> Path:
    run = {
        "experiment": args.experiment,
        "started_at": started,
        "runner_commit": git_commit(),
        "dataset": {"code": args.dataset, "path": str(args.data), "sha256": sha256(args.data), "row_count": len(df)},
        "env_lock_sha256": sha256(LOCK),
        "python": sys.version.split()[0],
        "seeds": seeds,
        "parts": {part: {"args": [Path(sys.argv[0]).name, json.dumps(args.params, ensure_ascii=False)],
                         "exit_code": exit_code, "seconds": round(time.time() - t0)}},
        "summary": summary,
        "results_path": str(out),
        "finished_at": datetime.now(UTC).isoformat(),
        "report": report,
    }
    path = out / "run.json"
    path.write_text(json.dumps(run, ensure_ascii=False, indent=2, default=float), encoding="utf-8")
    print(f"run.json: {path}")
    return path


def results_dir(root: Path) -> Path:
    out = root / datetime.now().strftime("%Y%m%d-%H%M%S")
    out.mkdir(parents=True)
    return out
